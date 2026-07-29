"""DOCX/XLSX/DOC/XLS conversion via an adapter architecture.

Backends are tried in the order configured in `config.toml`
(`[conversion] office_backend_order`), normally:

1. `libreoffice`  -- a local LibreOffice install (`soffice`), available
   and used for real conversions on this development platform as well
   as on Windows/macOS when LibreOffice is installed.
2. `office_com`   -- Microsoft Office automation via `pywin32`. Only
   ever attempted on Windows, and only if `pywin32` and Office are both
   installed. This adapter is exercised by the code path but has not
   been executed on a real Windows machine as part of this build; see
   the README's "Known Stage 1 Limitations" section.
3. `fallback`     -- a pure-Python best-effort renderer for modern
   `.docx` / `.xlsx` only (legacy `.doc` / `.xls` have no safe fallback
   and go through the standard unconverted-file procedure instead).

Whichever backend actually produced the output is recorded on the
result, and the fallback backend always adds an explicit warning so the
report never implies higher fidelity than was actually achieved.
"""

from __future__ import annotations

import platform
import shutil
import subprocess
import sys
import tempfile
import time
import uuid
from pathlib import Path
from xml.sax.saxutils import escape

from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..cancellation import CancellationToken
from ..models import ConversionOutcome, ConversionResult
from .base import failed_result, validate_pdf

_DOCX_EXT = {".docx"}
_XLSX_EXT = {".xlsx"}
_LEGACY_EXT = {".doc", ".xls"}

_LIBREOFFICE_TIMEOUT_SECONDS = 180

# How often to poll a running LibreOffice process for cancellation while
# waiting for it to finish -- see `_convert_with_libreoffice`'s docstring.
_CANCELLATION_POLL_INTERVAL_SECONDS = 0.2

_WINDOWS_LIBREOFFICE_CANDIDATES = [
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
]


def can_handle(extension: str) -> bool:
    return extension in _DOCX_EXT | _XLSX_EXT | _LEGACY_EXT


def convert(
    occurrence, dest_path: Path, config, workspace=None, cancellation_token: CancellationToken | None = None
) -> ConversionResult:
    source = occurrence.extracted_path
    if source is None or not source.exists():
        return failed_result("Original office document bytes were not available to convert.")

    extension = occurrence.original_extension
    backend_order = getattr(config, "office_backend_order", ("libreoffice", "office_com", "fallback"))

    attempts: list[str] = []
    for backend in backend_order:
        if backend == "libreoffice":
            soffice = _find_libreoffice()
            if not soffice:
                attempts.append("libreoffice: not found on this machine")
                continue
            result = _convert_with_libreoffice(soffice, source, dest_path, cancellation_token)
            if result is not None:
                return result
            attempts.append("libreoffice: conversion attempt failed")
        elif backend == "office_com":
            if sys.platform != "win32":
                attempts.append("office_com: not applicable (not running on Windows)")
                continue
            result = _convert_with_office_com(source, dest_path)
            if result is not None:
                return result
            attempts.append("office_com: not available or conversion failed")
        elif backend == "fallback":
            if extension in _LEGACY_EXT:
                attempts.append("fallback: no safe fallback exists for legacy .doc/.xls")
                continue
            result = _convert_with_fallback(occurrence, source, dest_path, extension)
            if result is not None:
                return result
            attempts.append("fallback: conversion attempt failed")

    detail = "; ".join(attempts) if attempts else "no conversion backend was configured"
    return failed_result(
        f"No available converter could produce a PDF for this {extension} file ({detail})."
    )


def _find_libreoffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    if platform.system() == "Windows":
        for candidate in _WINDOWS_LIBREOFFICE_CANDIDATES:
            if Path(candidate).exists():
                return candidate
    return None


def _convert_with_libreoffice(
    soffice: str, source: Path, dest_path: Path, cancellation_token: CancellationToken | None = None
) -> ConversionResult | None:
    """Runs LibreOffice headless conversion, polling for cancellation
    instead of blocking on a single `subprocess.run(..., timeout=...)`
    call.

    A real user report: pressing "Cancel Processing" while one of these
    was in flight left the app showing "Cancelling..." for minutes,
    because the old implementation used `subprocess.run` with a 180
    second timeout and no way to notice a cancellation request until
    that call returned. This polls the running process every
    `_CANCELLATION_POLL_INTERVAL_SECONDS` instead, so a cancellation
    request is noticed (and the process killed) almost immediately
    rather than only after this one file happens to finish or time out.
    """

    with tempfile.TemporaryDirectory(prefix="lpb_soffice_") as tmp_out:
        profile_dir = Path(tempfile.gettempdir()) / f"lpb_soffice_profile_{uuid.uuid4().hex}"
        cmd = [
            soffice,
            "--headless",
            "--norestore",
            "--nolockcheck",
            f"-env:UserInstallation=file:///{profile_dir.as_posix()}",
            "--convert-to",
            "pdf",
            "--outdir",
            tmp_out,
            str(source),
        ]
        try:
            proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        except OSError:
            shutil.rmtree(profile_dir, ignore_errors=True)
            return None

        try:
            deadline = time.monotonic() + _LIBREOFFICE_TIMEOUT_SECONDS
            while True:
                try:
                    proc.wait(timeout=_CANCELLATION_POLL_INTERVAL_SECONDS)
                    break
                except subprocess.TimeoutExpired:
                    if cancellation_token is not None and cancellation_token.is_requested():
                        proc.kill()
                        proc.wait()
                        return None
                    if time.monotonic() >= deadline:
                        proc.kill()
                        proc.wait()
                        return None
        finally:
            shutil.rmtree(profile_dir, ignore_errors=True)

        produced = Path(tmp_out) / (source.stem + ".pdf")
        if proc.returncode != 0 or not produced.exists():
            return None

        dest_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(produced, dest_path)

    is_valid, page_count, error = validate_pdf(dest_path)
    if not is_valid:
        return None

    return ConversionResult(
        outcome=ConversionOutcome.SUCCESS,
        pdf_path=dest_path,
        page_count=page_count,
        backend="libreoffice",
        warnings=[],
    )


def _convert_with_office_com(source: Path, dest_path: Path) -> ConversionResult | None:
    """Microsoft Office COM automation, Windows-only.

    This path requires `pywin32` and a licensed local Office install.
    It cannot be exercised on the Linux development machine used to
    build Stage 1; it is implemented defensively (any failure simply
    falls through to the next backend) rather than assumed to work.
    """

    try:
        import win32com.client  # type: ignore
    except ImportError:
        return None

    extension = source.suffix.lower()
    dest_path.parent.mkdir(parents=True, exist_ok=True)

    try:
        if extension in _DOCX_EXT | _LEGACY_EXT:
            app = win32com.client.DispatchEx("Word.Application")
            app.Visible = False
            try:
                doc = app.Documents.Open(str(source), ReadOnly=True)
                try:
                    doc.SaveAs(str(dest_path), FileFormat=17)  # wdFormatPDF
                finally:
                    doc.Close(False)
            finally:
                app.Quit()
        elif extension in _XLSX_EXT | _LEGACY_EXT:
            app = win32com.client.DispatchEx("Excel.Application")
            app.Visible = False
            try:
                wb = app.Workbooks.Open(str(source), ReadOnly=True)
                try:
                    wb.ExportAsFixedFormat(0, str(dest_path))  # xlTypePDF
                finally:
                    wb.Close(False)
            finally:
                app.Quit()
        else:
            return None
    except Exception:
        return None

    is_valid, page_count, error = validate_pdf(dest_path)
    if not is_valid:
        return None

    return ConversionResult(
        outcome=ConversionOutcome.SUCCESS,
        pdf_path=dest_path,
        page_count=page_count,
        backend="office_com",
        warnings=[],
    )


def _convert_with_fallback(occurrence, source: Path, dest_path: Path, extension: str) -> ConversionResult | None:
    if extension in _DOCX_EXT:
        return _fallback_docx(source, dest_path)
    if extension in _XLSX_EXT:
        return _fallback_xlsx(source, dest_path)
    return None


def _fallback_docx(source: Path, dest_path: Path) -> ConversionResult | None:
    try:
        import docx
    except ImportError:
        return None

    try:
        document = docx.Document(str(source))
    except Exception as exc:  # python-docx raises assorted errors on bad files
        return failed_result(f"Fallback DOCX reader could not open the file: {exc}")

    body_style = ParagraphStyle(name="Body", fontName="Helvetica", fontSize=10, leading=13)
    heading_style = ParagraphStyle(name="Heading", fontName="Helvetica-Bold", fontSize=13, leading=17)

    flowables = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = heading_style if para.style and "Heading" in (para.style.name or "") else body_style
        flowables.append(Paragraph(escape(text), style))
        flowables.append(Spacer(1, 4))

    for table in document.tables:
        rows = [[escape(cell.text) for cell in row.cells] for row in table.rows]
        if rows:
            t = Table(rows)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, "#666666"), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
            flowables.append(t)
            flowables.append(Spacer(1, 8))

    if not flowables:
        return failed_result("Fallback DOCX renderer found no readable paragraphs or tables.")

    dest_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(dest_path), pagesize=LETTER)
    try:
        doc.build(flowables)
    except Exception as exc:
        return failed_result(f"Fallback DOCX renderer failed to lay out content: {exc}")

    is_valid, page_count, error = validate_pdf(dest_path)
    if not is_valid:
        return failed_result(f"Fallback DOCX output failed validation: {error}")

    return ConversionResult(
        outcome=ConversionOutcome.FALLBACK_SUCCESS,
        pdf_path=dest_path,
        page_count=page_count,
        backend="fallback-docx",
        warnings=[
            "High-fidelity conversion (LibreOffice/Office) was not available. Used the basic "
            "fallback renderer: paragraph and table text only. Original formatting, images, "
            "headers/footers, and layout are NOT preserved."
        ],
    )


def _fallback_xlsx(source: Path, dest_path: Path) -> ConversionResult | None:
    try:
        import openpyxl
    except ImportError:
        return None

    try:
        wb = openpyxl.load_workbook(str(source), data_only=True, read_only=True)
    except Exception as exc:
        return failed_result(f"Fallback XLSX reader could not open the file: {exc}")

    heading_style = ParagraphStyle(name="Sheet", fontName="Helvetica-Bold", fontSize=13, leading=17)

    flowables = []
    any_content = False
    for i, sheet_name in enumerate(wb.sheetnames):
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            if row is None or all(v is None for v in row):
                continue
            rows.append([escape(str(v)) if v is not None else "" for v in row])
        if i > 0:
            flowables.append(PageBreak())
        flowables.append(Paragraph(f"Sheet: {escape(sheet_name)}", heading_style))
        flowables.append(Spacer(1, 6))
        if rows:
            any_content = True
            max_cols = max(len(r) for r in rows)
            rows = [r + [""] * (max_cols - len(r)) for r in rows]
            t = Table(rows[:2000])
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, "#888888"), ("FONTSIZE", (0, 0), (-1, -1), 7)]))
            flowables.append(t)
        else:
            flowables.append(Paragraph("(empty sheet)", ParagraphStyle(name="Empty", fontSize=9)))

    if not any_content:
        return failed_result("Fallback XLSX renderer found no non-empty cells in any worksheet.")

    dest_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(dest_path), pagesize=LETTER)
    try:
        doc.build(flowables)
    except Exception as exc:
        return failed_result(f"Fallback XLSX renderer failed to lay out content: {exc}")

    is_valid, page_count, error = validate_pdf(dest_path)
    if not is_valid:
        return failed_result(f"Fallback XLSX output failed validation: {error}")

    return ConversionResult(
        outcome=ConversionOutcome.FALLBACK_SUCCESS,
        pdf_path=dest_path,
        page_count=page_count,
        backend="fallback-xlsx",
        warnings=[
            "High-fidelity conversion (LibreOffice/Office) was not available. Used the basic "
            "fallback renderer: cell values only, as a plain grid. Original formatting, "
            "formulas, charts, and layout are NOT preserved."
        ],
    )
