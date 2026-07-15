"""Synthetic document builders used by the automated test suite.

Every fixture used in tests is generated programmatically here -- no
real borrower data is ever used or required.
"""

from __future__ import annotations

import email.policy
import random
import zipfile
from email.message import EmailMessage
from pathlib import Path

from pypdf import PdfReader, PdfWriter
from pypdf.generic import (
    ArrayObject,
    BooleanObject,
    DictionaryObject,
    FloatObject,
    NameObject,
    NumberObject,
    TextStringObject,
)
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


def make_pdf(path: Path, pages: int = 1, text_prefix: str = "Page", metadata: dict | None = None) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), pagesize=LETTER)
    for i in range(pages):
        c.setFont("Helvetica", 12)
        c.drawString(72, 700, f"{text_prefix} {i + 1} of {pages}")
        c.showPage()
    c.save()

    if metadata:
        reader = PdfReader(str(path))
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        writer.add_metadata(metadata)
        with path.open("wb") as fh:
            writer.write(fh)
    return path


def make_signature_package(path: Path, borrower_name: str, common_pages: int = 19) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), pagesize=LETTER)
    for i in range(common_pages):
        c.setFont("Helvetica", 12)
        c.drawString(72, 700, f"Disclosure page {i + 1} of {common_pages + 1} - shared content")
        c.showPage()
    c.setFont("Helvetica", 12)
    c.drawString(72, 700, f"Signature page - Borrower: {borrower_name}")
    c.showPage()
    c.save()
    return path


def make_pdf_with_blank_pages(
    path: Path, pages: int, blank_positions: list[int], text_prefix: str = "Page"
) -> Path:
    """Like `make_pdf`, but pages at `blank_positions` (0-indexed) are
    left genuinely blank instead of getting text -- used for RC2's
    blank-page-tolerant duplicate detection tests.
    """

    path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), pagesize=LETTER)
    for i in range(pages):
        if i not in blank_positions:
            c.setFont("Helvetica", 12)
            c.drawString(72, 700, f"{text_prefix} {i + 1} of {pages}")
        c.showPage()
    c.save()
    return path


def _draw_faint_mark(c: canvas.Canvas, x: int = 100, y: int = 400) -> None:
    """Draws a faint gray squiggle (signature-stroke-like) as a raster
    image, not vector text -- real wet signatures/stamps/marks on a
    scanned page are raster content, and a page containing one must
    never be classified blank no matter how faint it looks.
    """

    from PIL import Image, ImageDraw

    img = Image.new("RGB", (300, 150), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.line([(20, 100), (60, 60), (100, 110), (140, 70), (180, 100)], fill=(180, 180, 180), width=2)
    import io

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    c.drawImage(ImageReader(buf), x, y, width=300, height=150)


def make_pdf_with_faint_content(path: Path, blank_pages: int = 1, faint_page_index: int = 0) -> Path:
    """`blank_pages` total pages; the page at `faint_page_index` looks
    almost blank but carries a faint raster mark (a faint signature
    stroke) and must NOT be classified blank -- see TEST 6 in
    test_content_fingerprinting.py.
    """

    path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), pagesize=LETTER)
    for i in range(blank_pages):
        if i == faint_page_index:
            _draw_faint_mark(c)
        c.showPage()
    c.save()
    return path


def _add_text_field(writer: PdfWriter, page, name: str, value: str, rect: tuple[float, float, float, float]):
    widget = DictionaryObject()
    widget[NameObject("/FT")] = NameObject("/Tx")
    widget[NameObject("/Type")] = NameObject("/Annot")
    widget[NameObject("/Subtype")] = NameObject("/Widget")
    widget[NameObject("/T")] = TextStringObject(name)
    widget[NameObject("/V")] = TextStringObject(value)
    widget[NameObject("/Rect")] = ArrayObject([FloatObject(x) for x in rect])
    widget[NameObject("/F")] = NumberObject(4)
    widget_ref = writer._add_object(widget)
    widget[NameObject("/P")] = page.indirect_reference

    if "/Annots" in page:
        page["/Annots"].append(widget_ref)
    else:
        page[NameObject("/Annots")] = ArrayObject([widget_ref])

    root = writer._root_object
    if "/AcroForm" not in root:
        acroform = DictionaryObject()
        acroform[NameObject("/Fields")] = ArrayObject()
        acroform[NameObject("/NeedAppearances")] = BooleanObject(True)
        root[NameObject("/AcroForm")] = acroform
    root["/AcroForm"]["/Fields"].append(widget_ref)
    return widget_ref


def _add_signature_field(writer: PdfWriter, page, name: str, rect: tuple[float, float, float, float]):
    """A minimal (not cryptographically valid) /Sig field widget --
    sufficient to exercise "has a signed digital signature field"
    detection in pdf_content.py without needing a real PKI signing flow,
    which this fully-offline app never performs itself anyway (it only
    ever detects pre-existing signatures in supplied PDFs).
    """

    sig_value = DictionaryObject()
    sig_value[NameObject("/Type")] = NameObject("/Sig")
    sig_value[NameObject("/Filter")] = NameObject("/Adobe.PPKLite")
    sig_value[NameObject("/Name")] = TextStringObject(name)
    sig_value_ref = writer._add_object(sig_value)

    widget = DictionaryObject()
    widget[NameObject("/FT")] = NameObject("/Sig")
    widget[NameObject("/Type")] = NameObject("/Annot")
    widget[NameObject("/Subtype")] = NameObject("/Widget")
    widget[NameObject("/T")] = TextStringObject(name)
    widget[NameObject("/V")] = sig_value_ref
    widget[NameObject("/Rect")] = ArrayObject([FloatObject(x) for x in rect])
    widget[NameObject("/F")] = NumberObject(4)
    widget_ref = writer._add_object(widget)
    widget[NameObject("/P")] = page.indirect_reference

    if "/Annots" in page:
        page["/Annots"].append(widget_ref)
    else:
        page[NameObject("/Annots")] = ArrayObject([widget_ref])

    root = writer._root_object
    if "/AcroForm" not in root:
        acroform = DictionaryObject()
        acroform[NameObject("/Fields")] = ArrayObject()
        root[NameObject("/AcroForm")] = acroform
    root["/AcroForm"]["/Fields"].append(widget_ref)
    return widget_ref


def make_form_pdf(path: Path, field_values: dict[str, str], pages: int = 1, common_text: str = "Loan Application") -> Path:
    """A PDF with real AcroForm text-field widgets (name -> value), not
    just static text -- used for form-field-difference duplicate-
    detection tests (dates/conditions/checkboxes must never be silently
    ignored).
    """

    path.parent.mkdir(parents=True, exist_ok=True)
    buf_writer = PdfWriter()
    plain = path.with_suffix(".plain.pdf")
    make_pdf(plain, pages=pages, text_prefix=common_text)
    reader = PdfReader(str(plain))
    buf_writer.append_pages_from_reader(reader)

    y = 650
    for name, value in field_values.items():
        _add_text_field(buf_writer, buf_writer.pages[0], name, value, (100, y, 400, y + 20))
        y -= 30

    with path.open("wb") as fh:
        buf_writer.write(fh)
    plain.unlink()
    return path


def make_signed_pdf_variant(
    path: Path, borrower_name: str, common_pages: int = 19, signature_kind: str = "unsigned"
) -> Path:
    """Extends `make_signature_package`'s shared-boilerplate-plus-final-
    signature-page shape with a real distinguishing signal per
    `signature_kind`:
      - "unsigned":   final page has plain text only, no signature
                      apparatus at all.
      - "e_signed":   final page has a signed AcroForm /Sig field (a
                      digital e-signature) plus audit-style text.
      - "wet_signed": final page embeds a faint raster signature mark
                      (a scanned ink signature), no /Sig field at all --
                      wet signatures are physical marks, not AcroForm
                      fields.
    """

    if signature_kind not in ("unsigned", "e_signed", "wet_signed"):
        raise ValueError(f"unknown signature_kind: {signature_kind!r}")

    path.parent.mkdir(parents=True, exist_ok=True)
    plain = path.with_suffix(".plain.pdf")
    c = canvas.Canvas(str(plain), pagesize=LETTER)
    for i in range(common_pages):
        c.setFont("Helvetica", 12)
        c.drawString(72, 700, f"Disclosure page {i + 1} of {common_pages + 1} - shared content")
        c.showPage()

    c.setFont("Helvetica", 12)
    if signature_kind == "unsigned":
        c.drawString(72, 700, f"Signature page - Borrower: {borrower_name} - NOT YET SIGNED")
    elif signature_kind == "e_signed":
        c.drawString(72, 700, f"Signature page - Borrower: {borrower_name}")
        c.drawString(72, 680, "Electronically signed via e-signature audit trail")
    else:
        c.drawString(72, 700, f"Signature page - Borrower: {borrower_name}")
        _draw_faint_mark(c, x=72, y=550)
    c.showPage()
    c.save()

    if signature_kind == "e_signed":
        reader = PdfReader(str(plain))
        writer = PdfWriter()
        writer.append_pages_from_reader(reader)
        last_page = writer.pages[-1]
        _add_signature_field(writer, last_page, f"{borrower_name}_signature", (72, 600, 300, 620))
        with path.open("wb") as fh:
            writer.write(fh)
        plain.unlink()
    else:
        plain.rename(path)
    return path


def make_scanned_like_pdf(path: Path, pages: int, image_seed: int, size: tuple[int, int] = (850, 1100)) -> Path:
    """Image-only pages with NO extractable text layer at all --
    required to genuinely exercise the last-tier rendered-visual-
    comparison path (pdf_render.py): if a fixture has any real text,
    Level 3's text signal resolves ambiguity before the render tier is
    ever reached. Deterministic per `image_seed` (same seed -> same
    rendered content; different seed -> visibly different content),
    simulating scanned document pages.
    """

    from PIL import Image, ImageDraw
    import io

    path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), pagesize=LETTER)
    rng = random.Random(image_seed)
    for _ in range(pages):
        img = Image.new("RGB", size, color=(250, 250, 248))
        draw = ImageDraw.Draw(img)
        # A handful of deterministic pseudo-text-like gray blocks/lines,
        # standing in for a scanned page's visual content without using
        # any real extractable text.
        for _ in range(12):
            x0 = rng.randint(40, size[0] - 200)
            y0 = rng.randint(40, size[1] - 40)
            width = rng.randint(80, 180)
            gray = rng.randint(40, 90)
            draw.line([(x0, y0), (x0 + width, y0)], fill=(gray, gray, gray), width=3)
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        buf.seek(0)
        c.drawImage(ImageReader(buf), 0, 0, width=LETTER[0], height=LETTER[1])
        c.showPage()
    c.save()
    return path


def make_pdf_portfolio(
    path: Path,
    attachments: list[tuple[str, bytes]],
    cover_text: str = "Open this portfolio in Adobe Acrobat to view its contents.",
) -> Path:
    """A PDF with a /Collection dictionary on its root (a true Portfolio,
    per the /Collection-presence rule in pdf_portfolio.py) plus a cover
    page and a set of embedded attachments, in the given order.
    """

    path.parent.mkdir(parents=True, exist_ok=True)
    plain = path.with_suffix(".plain.pdf")
    c = canvas.Canvas(str(plain), pagesize=LETTER)
    c.setFont("Helvetica", 12)
    c.drawString(72, 700, cover_text)
    c.showPage()
    c.save()

    reader = PdfReader(str(plain))
    writer = PdfWriter()
    writer.append_pages_from_reader(reader)
    for name, data in attachments:
        writer.add_attachment(name, data)

    collection_dict = DictionaryObject()
    collection_dict[NameObject("/Type")] = NameObject("/Collection")
    collection_dict[NameObject("/View")] = NameObject("/D")
    writer._root_object[NameObject("/Collection")] = collection_dict

    with path.open("wb") as fh:
        writer.write(fh)
    plain.unlink()
    return path


def make_merged_pdf(path: Path, component_paths: list[Path]) -> Path:
    """Concatenates each component PDF's pages, in order, into one
    merged PDF -- a synthetic stand-in for a large merged lender
    package that already contains several smaller documents.
    """

    path.parent.mkdir(parents=True, exist_ok=True)
    writer = PdfWriter()
    for component in component_paths:
        reader = PdfReader(str(component))
        writer.append_pages_from_reader(reader)
    with path.open("wb") as fh:
        writer.write(fh)
    return path


def make_merged_pdf_with_gap(path: Path, component_paths: list[Path], extra_pages_between: list[Path]) -> Path:
    """Like `make_merged_pdf`, but inserts one page from
    `extra_pages_between` (cycled if shorter than needed) between each
    pair of consecutive components -- breaks contiguity so a component's
    pages are no longer a contiguous run inside the merged PDF, used for
    partial-overlap / non-containment tests.
    """

    path.parent.mkdir(parents=True, exist_ok=True)
    writer = PdfWriter()
    for i, component in enumerate(component_paths):
        reader = PdfReader(str(component))
        writer.append_pages_from_reader(reader)
        if i < len(component_paths) - 1 and extra_pages_between:
            gap_reader = PdfReader(str(extra_pages_between[i % len(extra_pages_between)]))
            writer.append_pages_from_reader(gap_reader)
    with path.open("wb") as fh:
        writer.write(fh)
    return path


def make_corrupt_pdf(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"%PDF-1.4\nThis is not a valid PDF body.\n%%EOF-BUT-NOT-REALLY")
    return path


def make_password_protected_pdf(path: Path, password: str = "secret123", pages: int = 2) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    plain = path.with_suffix(".plain.pdf")
    make_pdf(plain, pages=pages, text_prefix="Protected page")
    reader = PdfReader(str(plain))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt(user_password=password, owner_password=None)
    with path.open("wb") as fh:
        writer.write(fh)
    plain.unlink()
    return path


def make_image(path: Path, fmt: str = "PNG", size: tuple[int, int] = (300, 200), color=(200, 30, 30)) -> Path:
    from PIL import Image

    path.parent.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", size, color)
    img.save(path, format=fmt)
    return path


def make_multipage_tiff(path: Path, frames: int = 3, size: tuple[int, int] = (200, 150)) -> Path:
    from PIL import Image

    path.parent.mkdir(parents=True, exist_ok=True)
    images = [Image.new("RGB", size, (10 * i, 100, 200 - 10 * i)) for i in range(frames)]
    images[0].save(path, format="TIFF", save_all=True, append_images=images[1:])
    return path


def make_txt(path: Path, content: str, encoding: str = "utf-8") -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content.encode(encoding))
    return path


def make_html(path: Path, body_html: str, title: str = "Test Document") -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    content = f"""<!DOCTYPE html>
<html><head><title>{title}</title>
<script>fetch('https://example.com/track');</script>
</head><body>{body_html}</body></html>"""
    path.write_text(content, encoding="utf-8")
    return path


def make_docx(path: Path, paragraphs: list[str], table_rows: list[list[str]] | None = None) -> Path:
    import docx

    path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    document.add_heading(paragraphs[0] if paragraphs else "Document", level=1)
    for para in paragraphs[1:]:
        document.add_paragraph(para)
    if table_rows:
        table = document.add_table(rows=0, cols=len(table_rows[0]))
        for row in table_rows:
            cells = table.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = value
    document.save(str(path))
    return path


def make_xlsx(path: Path, rows: list[list[str]], sheet_name: str = "Sheet1") -> Path:
    import openpyxl

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    wb.save(str(path))
    return path


def make_eml(
    path: Path,
    subject: str = "Test message",
    from_addr: str = "sender@example.com",
    to_addr: str = "recipient@example.com",
    body: str = "This is the message body.",
    attachments: list[tuple[str, bytes, str, str]] | None = None,
) -> Path:
    """attachments: list of (filename, data, maintype, subtype)."""

    path.parent.mkdir(parents=True, exist_ok=True)
    msg = EmailMessage(policy=email.policy.default)
    msg["Subject"] = subject
    msg["From"] = from_addr
    msg["To"] = to_addr
    msg.set_content(body)

    for filename, data, maintype, subtype in attachments or []:
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=filename)

    path.write_bytes(msg.as_bytes())
    return path


def make_zip(path: Path, entries: list[tuple[str, bytes]]) -> Path:
    """entries preserved in the exact order given."""

    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries:
            zf.writestr(name, data)
    return path


def read_pdf_page_count(path: Path) -> int:
    return len(PdfReader(str(path)).pages)


def read_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)
